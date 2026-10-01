"""Generate corrected Excel and audit report."""

from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from utils import ALLERGENS, safe_str, to_int


LIGHT_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")


def _find_allergen_col(cols, allergen):
    for col in cols:
        if allergen.lower() in str(col).lower():
            return col
    return None


def _find_col(cols, keywords):
    for kw in keywords:
        kw_l = kw.lower()
        for col in cols:
            if kw_l in str(col).lower():
                return col
    return None


def generate_corrected_excel(file_bytes, rm_sheet, sfg_sheet, sku_sheet,
                             rm_map, sfg_expected, sku_expected,
                             rm_code_col, sfg_code_col, sku_code_col):
    """Load the original workbook with openpyxl, fix allergen values, highlight in green."""
    wb = load_workbook(BytesIO(file_bytes))

    def fix_sheet(sheet_name, code_col, expected_map):
        if sheet_name is None or code_col is None:
            return
        ws = wb[sheet_name]
        # Find header row (row 1) and map column index
        header_row = {}
        for cell in ws[1]:
            if cell.value is not None:
                header_row[str(cell.value).strip()] = cell.column

        # Find code column index
        code_idx = None
        for h, idx in header_row.items():
            if code_col and h.lower() == str(code_col).lower():
                code_idx = idx
                break
            if "code" in h.lower():
                code_idx = idx
                break
        if code_idx is None:
            return

        # Find allergen column indices
        allergen_idx = {}
        for allergen in ALLERGENS:
            for h, idx in header_row.items():
                if allergen.lower() in h.lower():
                    allergen_idx[allergen] = idx
                    break

        for row in ws.iter_rows(min_row=2):
            code_cell = row[code_idx - 1]
            code = safe_str(code_cell.value)
            if code not in expected_map:
                continue
            exp = expected_map[code]
            for allergen, col_idx in allergen_idx.items():
                exp_val = exp.get(allergen, 0)
                cell = row[col_idx - 1]
                cur = to_int(cell.value)
                if cur != exp_val:
                    cell.value = exp_val
                    cell.fill = LIGHT_GREEN

    fix_sheet(rm_sheet, rm_code_col, rm_map)
    fix_sheet(sfg_sheet, sfg_code_col, sfg_expected)
    fix_sheet(sku_sheet, sku_code_col, sku_expected)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    wb.close()
    return out.getvalue()


def generate_audit_report(errors, summary):
    """Generate an audit report workbook with summary + error detail."""
    out = BytesIO()
    with pd.ExcelWriter(out, engine="xlsxwriter") as writer:
        # Summary sheet
        summary_df = pd.DataFrame([
            {"Metric": "RM Checked", "Value": summary.get("rm_checked", 0)},
            {"Metric": "SFG Allergen Mismatches", "Value": summary.get("sfg_mismatches", 0)},
            {"Metric": "SKU Checked", "Value": summary.get("sku_checked", 0)},
            {"Metric": "Formula Errors", "Value": summary.get("formula_errors", 0)},
            {"Metric": "Total Errors", "Value": len(errors)},
        ])
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

        # Errors sheet
        if errors:
            err_df = pd.DataFrame([e.__dict__ if hasattr(e, "__dict__") else e for e in errors])
        else:
            err_df = pd.DataFrame(columns=["Sheet", "Row", "Type", "Expected", "Found", "Message"])
        err_df.to_excel(writer, sheet_name="Audit Details", index=False)

        # Format
        workbook = writer.book
        header_fmt = workbook.add_format({
            "bold": True, "bg_color": "#4472C4", "font_color": "white",
            "border": 1, "text_wrap": True,
        })
        cell_fmt = workbook.add_format({"text_wrap": True, "border": 1, "valign": "top"})
        for sheet_name in ["Summary", "Audit Details"]:
            ws = writer.sheets[sheet_name]
            ws.set_column(0, 5, 22, cell_fmt)
            for col in range(6):
                ws.write(0, col, "", header_fmt)

    out.seek(0)
    return out.getvalue()
