"""Formula inspection rules using openpyxl.

Opens the uploaded workbook with openpyxl (data_only=False) to read formulas.
Flags XLOOKUP/VLOOKUP that reference a Name column instead of a Code column.
"""

from openpyxl import load_workbook
from io import BytesIO
import re


def verify_formulas(file_bytes, result):
    """Scan all sheets for VLOOKUP/XLOOKUP formulas referencing Name columns."""
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=False)
    except Exception as e:
        result.add("Workbook", 0, "Formula", "Readable workbook", str(e),
                   f"Could not open workbook for formula inspection: {e}")
        return

    lookup_pattern = re.compile(r"(XLOOKUP|VLOOKUP)\s*\(", re.IGNORECASE)

    for ws in wb.worksheets:
        # Build a map of column letter -> header name from row 1
        headers = {}
        for cell in ws[1]:
            if cell.value is not None:
                headers[cell.column_letter] = str(cell.value).lower()

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value is None:
                    continue
                val = str(cell.value)
                if not lookup_pattern.search(val):
                    continue

                # Look for column references like B:B or Table1[Name] or A:B
                # Flag if formula references a Name column
                # Check for [Name] or [name] table references
                if re.search(r"\[name\]", val, re.IGNORECASE):
                    result.add(
                        ws.title, cell.row, "Formula",
                        "Lookup on Code column", "Lookup on Name column",
                        f"{cell.coordinate}: formula uses Name column for lookup: {val}",
                    )
                    result.formula_errors += 1
                    continue

                # Check range references like B2:D100 — identify if the lookup col is a Name col
                # Extract the first column letter from ranges like VLOOKUP(A2, B:D, ...)
                range_match = re.search(r"(XLOOKUP|VLOOKUP)\s*\([^,]+,\s*([A-Z]+)\d*:[A-Z]+", val, re.IGNORECASE)
                if range_match:
                    col_letter = range_match.group(2).upper()
                    header = headers.get(col_letter, "")
                    if "name" in header and "code" not in header:
                        result.add(
                            ws.title, cell.row, "Formula",
                            "Lookup on Code column", f"Lookup on {header} column",
                            f"{cell.coordinate}: VLOOKUP/XLOOKUP should use Code column, found '{header}'",
                        )
                        result.formula_errors += 1

    wb.close()
